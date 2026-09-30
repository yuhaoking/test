import json
import os
import queue
import shutil
import subprocess
import sys
import threading
import time
import wave

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
tmp = os.path.join(ROOT, '.tmp-office')
os.makedirs(tmp, exist_ok=True)
RESULTS = []


def mark(name, ok, extra=''):
    RESULTS.append((name, ok))
    print(('PASS ' if ok else 'FAIL ') + name + (' | ' + str(extra) if extra else ''))


class Rpc:
    def __init__(self):
        self.child = subprocess.Popen(
            [sys.executable, os.path.join(ROOT, 'plugins', 'office-convert-all', 'main.py')],
            cwd=ROOT,
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            text=True,
            encoding='utf-8',
            env={**os.environ, 'PYTHONUTF8': '1', 'PLUGIN_SAVE_DIR': tmp}
        )
        self.rid = 0
        self.q = queue.Queue()
        threading.Thread(target=self._reader, daemon=True).start()

    def _reader(self):
        for line in self.child.stdout:
            line = line.strip()
            if not line:
                continue
            try:
                self.q.put(json.loads(line))
            except Exception:
                pass

    def call(self, method, params, timeout=300):
        self.rid += 1
        rid = self.rid
        self.child.stdin.write(json.dumps({'jsonrpc': '2.0', 'id': rid, 'method': method, 'params': params}) + '\n')
        self.child.stdin.flush()
        deadline = time.time() + timeout
        while time.time() < deadline:
            try:
                msg = self.q.get(timeout=0.2)
            except queue.Empty:
                continue
            if msg.get('id') == rid:
                return msg
        return None

    def close(self):
        try:
            self.child.kill()
        except Exception:
            pass


def conv(rpc, path, target, extra=None):
    values = {'paths': path, 'target': target}
    if extra:
        values.update(extra)
    return rpc.call('plugin.handle_action', {'action': 'convert', 'values': values})


def message(res):
    return (res or {}).get('result', {}).get('message', 'NO-RESULT')


def last_file(suffix, prefix=None):
    cands = [f for f in os.listdir(tmp) if f.endswith(suffix)]
    if prefix:
        cands = [f for f in cands if f.startswith(prefix)]
    if not cands:
        return None
    return os.path.join(tmp, sorted(cands)[-1])


# ---------- samples ----------
sys.path.insert(0, os.path.join(ROOT, 'plugins', '_shared'))
import screen_capture as sc

img1 = os.path.join(tmp, 'sample1.png')
img2 = os.path.join(tmp, 'sample2.png')
data, w, h = sc.capture_full()
sc.save_png(img1, w, h, data)
sc.save_png(img2, w, h, data)

wavfile = os.path.join(tmp, 'tone.wav')
with wave.open(wavfile, 'wb') as f:
    f.setnchannels(1)
    f.setsampwidth(2)
    f.setframerate(16000)
    for i in range(16000):
        f.writeframes((int(12000 * (i % 40 < 20) - 12000)).to_bytes(2, 'little', signed=True))

txt_utf8 = os.path.join(tmp, 'note.txt')
with open(txt_utf8, 'w', encoding='utf-8') as f:
    f.write('小鹏工具箱转换测试')
txt_gbk = os.path.join(tmp, 'note_gbk.txt')
with open(txt_gbk, 'wb') as f:
    f.write('你好世界'.encode('gbk'))

md = os.path.join(tmp, 'doc.md')
with open(md, 'w', encoding='utf-8') as f:
    f.write('# 标题\n\n正文内容，**加粗** 与 `code`。\n\n- 项目1\n- 项目2\n')

# pptx/docx/xlsx via python libs
from pptx import Presentation
p = Presentation()
for i in range(1, 3):
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = '测试幻灯片 %d' % i
    s.placeholders[1].text = '内容段落 %d' % i
pptx = os.path.join(tmp, 'demo.pptx')
p.save(pptx)

from docx import Document
d = Document()
d.add_heading('测试文档', 0)
d.add_paragraph('小鹏工具箱 Office 转换测试段落。')
docx_f = os.path.join(tmp, 'demo.docx')
d.save(docx_f)

import openpyxl
wb = openpyxl.Workbook()
ws = wb.active
ws['A1'] = '名称'
ws['B1'] = '数值'
ws.append(['测试', 42])
ws.append(['小鹏', 100])
xlsx = os.path.join(tmp, 'demo.xlsx')
wb.save(xlsx)

# video sample: 2s mp4 via mss+imageio
vid = os.path.join(tmp, 'demo_video.mp4')
try:
    import mss
    import numpy as np
    import imageio.v2 as imageio
    writer = imageio.get_writer(vid, fps=10, codec='libx264', quality=5)
    deadline = time.time() + 2
    with mss.mss() as sct:
        while time.time() < deadline:
            writer.append_data(np.asarray(sct.grab(sct.monitors[1]))[:, :, :3][:, :, ::-1])
    writer.close()
except Exception as e:
    print('video sample fail:', e)

rpc = Rpc()
try:
    ui = rpc.call('plugin.get_ui', {'values': {'paths': ''}})
    text = (ui or {}).get('result', {}).get('text', '')
    mark('engine hints shown', '工作文件夹引擎' in text, text.split('\n')[-1])

    # 图片
    for tgt in ['jpg', 'bmp', 'webp', 'gif', 'tiff']:
        res = conv(rpc, img1, tgt)
        msg = message(res)
        file = last_file('.' + tgt, 'sample1_to')
        mark('png->%s' % tgt, '转换完成' in msg and file and os.path.getsize(file) > 10000, msg[:70])
    res = conv(rpc, img1, 'ico')
    file = last_file('.ico', 'sample1_to')
    mark('png->ico', '转换完成' in message(res) and file, message(res)[:70])
    res = conv(rpc, img1 + ';' + img2, 'pdf')
    file = last_file('_merged.pdf')
    mark('images->merged pdf', '转换完成' in message(res) and file, message(res)[:70])

    # 视频
    if os.path.exists(vid):
        for tgt in ['mov', 'mkv', 'webm']:
            res = conv(rpc, vid, tgt)
            msg = message(res)
            ok = ('转换完成' in msg) and os.path.exists(os.path.join(tmp, 'demo_video_to.' + tgt))
            mark('mp4->%s' % tgt, ok, msg[:70])
        # 带音轨版本用于提取测试
        vid_audio = os.path.join(tmp, 'demo_audio.mp4')
        subprocess.run([os.path.join(ROOT, 'engines', 'bin', 'ffmpeg.exe'), '-y', '-f', 'lavfi', '-i',
                        'sine=frequency=440:duration=2', '-i', vid, '-c:v', 'copy', '-c:a', 'aac',
                        '-shortest', vid_audio], capture_output=True, timeout=120)
        for tgt in ['mp3', 'wav']:
            res = conv(rpc, vid_audio, tgt)
            msg = message(res)
            of = os.path.join(tmp, 'demo_audio_to.' + tgt)
            ok = '转换完成' in msg and os.path.exists(of) and os.path.getsize(of) > 1000
            mark('mp4(audio)>%s' % tgt, ok, msg[:70])
    else:
        mark('video sample', False, 'sample missing')

    # 音频
    for tgt in ['mp3', 'flac', 'ogg']:
        res = conv(rpc, wavfile, tgt)
        msg = message(res)
        ok = '转换完成' in msg and os.path.exists(os.path.join(tmp, 'tone_to.' + tgt))
        mark('wav->%s' % tgt, ok, msg[:70])

    # 文档 via LibreOffice
    for src, name in [(pptx, 'pptx'), (docx_f, 'docx'), (xlsx, 'xlsx')]:
        res = conv(rpc, src, 'pdf')
        msg = message(res)
        base = os.path.splitext(os.path.basename(src))[0]
        out_pdf = os.path.join(tmp, base + '.pdf') if os.path.exists(os.path.join(tmp, base + '.pdf')) else os.path.join(tmp, base + '_to.pdf')
        mark(name + '->pdf', '转换完成' in msg and os.path.exists(out_pdf) and os.path.getsize(out_pdf) > 10000, msg[:80])

    # pandoc: md->docx / docx->md
    res = conv(rpc, md, 'docx')
    msg = message(res)
    mark('md->docx(pandoc)', '转换完成' in msg and os.path.exists(os.path.join(tmp, 'doc_to.docx')), msg[:80])
    res2 = conv(rpc, docx_f, 'md')
    msg2 = message(res2)
    mark('docx->md(pandoc)', '转换完成' in msg2 and os.path.exists(os.path.join(tmp, 'demo_to.md')), msg2[:80])

    # PDF 操作（用 pymupdf 造一个带文字的 pdf）
    pdfsrc = os.path.join(tmp, 'text.pdf')
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), 'Hello Xiaopeng 12345')
    page2 = doc.new_page()
    page2.insert_text((72, 72), 'Second page text')
    doc.save(pdfsrc)
    doc.close()
    res = conv(rpc, pdfsrc, 'pdf-text')
    msg = message(res)
    txtfile = os.path.join(tmp, 'text_to.txt')
    content = ''
    if os.path.exists(txtfile):
        content = open(txtfile, encoding='utf-8').read()
    mark('pdf->text', '转换完成' in msg and 'Xiaopeng' in content, msg[:60] + ' | ' + content[:30])
    res = conv(rpc, pdfsrc, 'pdf-split', {'pages': '2'})
    msg = message(res)
    mark('pdf->split p2', '转换完成' in msg and os.path.exists(os.path.join(tmp, 'text_p1.pdf')), msg[:70])
    res = conv(rpc, pdfsrc, 'pdf-encrypt', {'password': '123'})
    msg = message(res)
    enc = os.path.join(tmp, 'text_encrypted.pdf')
    mark('pdf->encrypt', '转换完成' in msg and os.path.exists(enc), msg[:70])
    res = conv(rpc, enc, 'pdf-decrypt', {'password': '123'})
    mark('pdf->decrypt', '转换完成' in message(res) and os.path.exists(os.path.join(tmp, 'text_encrypted_decrypted.pdf')), message(res)[:70])
    res = conv(rpc, pdfsrc, 'png-pages')
    msg = message(res)
    mark('pdf->images(pymupdf)', '转换完成' in msg and os.path.exists(os.path.join(tmp, 'text_p01.png')), msg[:70])

    # 编码
    res = conv(rpc, txt_gbk, 'utf-8')
    msg = message(res)
    out = os.path.join(tmp, 'note_gbk_utf-8.txt')
    content = open(out, encoding='utf-8').read() if os.path.exists(out) else ''
    mark('gbk->utf-8', '转换完成' in msg and '你好世界' == content, msg[:60])
finally:
    rpc.close()

fails = [r for r in RESULTS if not r[1]]
print('-' * 40)
print('TOTAL: %d  PASS: %d  FAIL: %d' % (len(RESULTS), len(RESULTS) - len(fails), len(fails)))
if fails:
    print('FAILED:', [f[0] for f in fails])
